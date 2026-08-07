import asyncio
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches
import os

async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        # Use 1333x750 to match exactly 16:9 ratio for pptx (13.333 x 7.5 inches * 100 dpi)
        page = await browser.new_page(viewport={"width": 1333, "height": 750})
        
        file_path = f"file://{os.path.abspath('presentation/index.html')}"
        await page.goto(file_path)
        
        # Wait for fonts and Math (KaTeX) to render
        await page.wait_for_timeout(2000)
        
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        total_slides = 28
        
        for i in range(total_slides):
            # Wait for slide transitions
            await page.wait_for_timeout(600)
            
            # On slides 8 and 12 (0-indexed 7 and 11), advance the internal steps to show the final state
            if i == 7 or i == 11:
                for _ in range(3):
                    await page.keyboard.press("ArrowRight")
                    await page.wait_for_timeout(600)
            
            # Take screenshot
            screenshot_path = f"/tmp/slide_{i}.png"
            await page.screenshot(path=screenshot_path)
            
            # Add to PPTX
            blank_slide_layout = prs.slide_layouts[6] 
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(screenshot_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            
            # Go to next slide
            await page.keyboard.press("ArrowRight")

        # Save PPTX to presentation folder
        output_path = os.path.abspath('presentation/Apresentacao_Spectrum.pptx')
        prs.save(output_path)
        await browser.close()
        print(f"Done! PPTX saved at {output_path}")

asyncio.run(main())
